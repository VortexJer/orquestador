"""Generacion de .docx: lo que se fija son los dos fallos vistos en vivo -
el nivel de heading que llegaba como str y reventaba, y el markdown
(**negrita**) que salia con los asteriscos a la vista en vez de en negrita."""
from docx import Document

from app.tools.docx_tool import _nivel_valido, create_docx


def _texto_todo(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_nivel_de_heading_str_no_revienta(tmp_path):
    """El modelo manda level como '2' (str); python-docx hace 0<=level<=9 y
    petaba con 'int vs str'. Se coacciona."""
    assert _nivel_valido("2") == 2
    assert _nivel_valido(1.0) == 1
    assert _nivel_valido("basura") == 1
    assert _nivel_valido(99) == 9
    out = create_docx([{"heading": "Motor", "level": "2", "text": "B47"}], tmp_path / "x.docx")
    assert out.exists()


def test_markdown_negrita_se_convierte_en_formato_real(tmp_path):
    out = create_docx(
        [{"text": "El M5 monta un **V8 biturbo** y *xDrive*."}], tmp_path / "md.docx"
    )
    doc = Document(out)
    # No quedan asteriscos literales a la vista
    assert "**" not in _texto_todo(doc)
    assert "*" not in _texto_todo(doc)
    # y 'V8 biturbo' quedo en negrita real, 'xDrive' en cursiva
    runs = {r.text: (r.bold, r.italic) for p in doc.paragraphs for r in p.runs}
    assert runs.get("V8 biturbo") == (True, False)
    assert runs.get("xDrive") == (False, True)


def test_markdown_en_heading_y_tabla_se_limpia(tmp_path):
    out = create_docx([
        {"heading": "El **motor** S63", "level": 1},
        {"table": {"headers": ["Dato", "**Valor**"], "rows": [["Potencia", "**625 CV**"]]}},
    ], tmp_path / "ht.docx")
    doc = Document(out)
    assert "El motor S63" in _texto_todo(doc)
    assert "**" not in _texto_todo(doc)
    celdas = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert "Valor" in celdas and "625 CV" in celdas
    assert not any("*" in c for c in celdas)


def test_listas_con_markdown_conservan_su_estilo(tmp_path):
    out = create_docx(
        [{"list_items": ["Pros: **625 CV**", "Cons: *consumo*"], "ordered": False}],
        tmp_path / "li.docx",
    )
    doc = Document(out)
    parrs = [p for p in doc.paragraphs if p.text.strip()]
    assert all(p.style.name == "List Bullet" for p in parrs)
    assert "**" not in _texto_todo(doc)


def test_forma_de_bloque_del_skill_se_normaliza(tmp_path):
    """El skill office-word ensena bloques con `type` (`heading`/`paragraph`/
    `list` con `items`/`table` con `header`), forma que NO coincide con la que
    consumia create_docx (`heading`/`list_items`/`table.headers`). Antes esto
    salia roto: heading colado como parrafo, lista y tabla PERDIDAS. Ahora el
    normalizador acepta esa forma y el .docx sale entero."""
    out = create_docx([
        {"type": "heading", "level": 1, "text": "Historia del cafe"},
        {"type": "paragraph", "text": "El cafe es una bebida muy consumida."},
        {"type": "list", "items": ["Moler el grano", "Calentar el agua"]},
        {"type": "table", "header": ["Variedad", "Perfil"],
         "rows": [["Arabica", "Suave"], ["Robusta", "Amargo"]]},
    ], tmp_path / "bloques.docx")
    doc = Document(out)
    # El heading es un Heading REAL, no un parrafo suelto.
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Historia del cafe" in headings
    # La lista existe con estilo nativo (no se perdio).
    listas = [p.text for p in doc.paragraphs if p.style.name.startswith("List")]
    assert any("Moler el grano" in t for t in listas)
    # La tabla existe con su encabezado (no se perdio) y mapeo header->headers.
    celdas = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert "Variedad" in celdas and "Arabica" in celdas
