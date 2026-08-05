"""Tests de groq_agent que no requieren OPENROUTER_API_KEY: ejercitan
ToolExecutor y el confinamiento de rutas directamente, sin llamar a la
API del proveedor."""
import pytest

from groq_agent.safety import PathEscapeError, resolve_within_root
from groq_agent.skills_loader import build_system_prompt, list_available_skills
from groq_agent.tools import ToolExecutor


@pytest.fixture
def executor(tmp_path):
    return ToolExecutor(tmp_path)


def test_write_then_read_file(executor):
    executor.write_file("a.py", "x = 1\n")
    assert executor.read_file("a.py") == "x = 1\n"


def test_edit_file_requires_unique_match(executor):
    executor.write_file("a.py", "x = 1\nx = 1\n")
    result = executor.edit_file("a.py", "x = 1", "x = 2")
    assert "ERROR" in result and "2 veces" in result


def test_edit_file_applies_unique_match(executor):
    executor.write_file("a.py", "def f():\n    return 1\n")
    result = executor.edit_file("a.py", "return 1", "return 2")
    assert "Editado" in result
    assert "return 2" in executor.read_file("a.py")


def test_path_escape_blocked(tmp_path):
    with pytest.raises(PathEscapeError):
        resolve_within_root(tmp_path, "../../etc/passwd")


def test_ruta_relativa_se_queda_en_el_workspace(tmp_path):
    from groq_agent.safety import resolve_write_destination

    destino, fuera = resolve_write_destination(tmp_path, "informe.docx")
    assert fuera is False
    assert str(destino).startswith(str(tmp_path.resolve()))


def test_carpeta_conocida_resuelve_fuera(tmp_path):
    from pathlib import Path

    from groq_agent.safety import resolve_write_destination

    destino, fuera = resolve_write_destination(tmp_path, "descargas/x.docx")
    assert fuera is True
    assert destino == (Path.home() / "Downloads" / "x.docx").resolve() or "Downloads" in str(destino)


def test_ruta_absoluta_resuelve_fuera(tmp_path):
    from pathlib import Path

    from groq_agent.safety import resolve_write_destination

    fuera_dir = Path.home() / "algo_de_prueba" / "a.txt"
    destino, fuera = resolve_write_destination(tmp_path, str(fuera_dir))
    assert fuera is True


def test_relativa_con_traversal_sigue_bloqueada(tmp_path):
    from groq_agent.safety import resolve_write_destination

    with pytest.raises(PathEscapeError):
        resolve_write_destination(tmp_path, "../../fuera.txt")


def test_escribir_fuera_con_yes_se_rechaza(tmp_path):
    # auto_yes = sesion desatendida: no hay quien confirme, asi que escribir
    # fuera del workspace se rechaza en vez de hacerlo a ciegas.
    ex = ToolExecutor(tmp_path, auto_yes=True)
    res = ex.write_file("descargas/prueba_orq_test.txt", "hola")
    assert res.startswith("ERROR")
    assert "workspace" in res


def test_escribir_relativo_con_yes_va_al_workspace(tmp_path):
    ex = ToolExecutor(tmp_path, auto_yes=True)
    res = ex.write_file("dentro.txt", "hola")
    assert res.startswith("Escrito")
    assert (tmp_path / "dentro.txt").read_text() == "hola"


def test_escribir_fuera_sin_aprobar_se_rechaza(tmp_path):
    # Interactivo (sin --yes) pero SIN aprobacion previa de agent_loop: la
    # tool no escribe (la confirmacion la hace agent_loop en el hilo
    # principal, no la tool).
    ex = ToolExecutor(tmp_path, auto_yes=False)
    res = ex.write_file("descargas/prueba_orq_no.txt", "x")
    assert res.startswith("ERROR")
    assert "confirmacion" in res


def test_carpeta_externa_aprobada_permite_escribir(tmp_path, tmp_path_factory):
    # Simula lo que hace agent_loop tras un 'Si, y todo en esta carpeta':
    # aprueba la carpeta y entonces la tool escribe sin volver a exigir OK.
    fuera = tmp_path_factory.mktemp("fuera_del_ws")
    ex = ToolExecutor(tmp_path, auto_yes=False)
    ex.aprobar_carpeta_externa(fuera)
    destino = fuera / "sub" / "a.txt"
    res = ex.write_file(str(destino), "hola")
    assert res.startswith("Escrito")
    assert destino.read_text() == "hola"


def test_carpeta_aprobada_se_puede_LEER_no_solo_escribir(tmp_path, tmp_path_factory):
    # El bug de la web en Descargas: tras escribirla fuera hay que poder
    # RE-LEERLA/verificarla ahi, no solo escribirla. La carpeta aprobada es
    # area de trabajo completa.
    fuera = tmp_path_factory.mktemp("web_ext")
    ex = ToolExecutor(tmp_path, auto_yes=False)
    ex.aprobar_carpeta_externa(fuera)
    (fuera / "index.html").write_text("<h1>hola</h1>", encoding="utf-8")
    leido = ex.read_file(str(fuera / "index.html"))
    assert "hola" in leido
    lint = ex.lint_web_page(str(fuera / "index.html"))
    assert "ERROR" not in lint


def test_editar_en_carpeta_aprobada_va_a_esa_carpeta(tmp_path, tmp_path_factory):
    # El bug del transcript: edit_file mostraba Downloads pero editaba en
    # workspace\descargas. Ahora una edicion en carpeta aprobada va ahi.
    fuera = tmp_path_factory.mktemp("web_edit")
    ex = ToolExecutor(tmp_path, auto_yes=False)
    ex.aprobar_carpeta_externa(fuera)
    (fuera / "styles.css").write_text("body{color:red}", encoding="utf-8")
    res = ex.edit_file(str(fuera / "styles.css"), "red", "blue")
    assert res.startswith("Editado")
    assert (fuera / "styles.css").read_text() == "body{color:blue}"
    # y NADA en el workspace
    assert not (tmp_path / "styles.css").exists()


def test_archivos_tocados_registra_ruta_absoluta(tmp_path):
    ex = ToolExecutor(tmp_path)
    ex.write_file("index.html", "<h1>x</h1>")
    ex.edit_file("index.html", "x", "y")
    rutas = list(ex.archivos_tocados)
    assert len(rutas) == 1  # mismo archivo, no duplicado
    assert rutas[0].endswith("index.html")
    assert ex.archivos_tocados[rutas[0]] == "editado"  # ultima accion gana


def test_refrescar_manifiesto_va_al_final_y_no_se_duplica(tmp_path):
    from groq_agent import agent_loop

    ex = ToolExecutor(tmp_path)
    ex.write_file("a.txt", "1")
    ex.write_file("b.txt", "2")
    messages = [{"role": "system", "content": "sys"}, {"role": "user", "content": "tarea"}]
    agent_loop._refrescar_manifiesto(messages, ex)
    agent_loop._refrescar_manifiesto(messages, ex)  # dos veces: no debe acumular
    manifiestos = [m for m in messages if m["content"].startswith(agent_loop._MARCA_MANIFIESTO)]
    assert len(manifiestos) == 1
    assert messages[-1] is manifiestos[0]  # siempre al final
    assert "a.txt" in messages[-1]["content"] and "b.txt" in messages[-1]["content"]


def test_refrescar_plan_mantiene_el_checklist_vivo_al_final(tmp_path):
    """El checklist se reinyecta al final del contexto cada turno (como el
    manifiesto) para que el modelo no lo pierda de vista y lo vaya marcando."""
    from groq_agent import agent_loop

    ex = ToolExecutor(tmp_path)
    ex.plan([
        {"paso": "Investigar", "estado": "hecho"},
        {"paso": "Escribir el HTML", "estado": "en_curso"},
        {"paso": "Verificar", "estado": "pendiente"},
    ])
    messages = [{"role": "system", "content": "sys"}, {"role": "user", "content": "tarea"}]
    agent_loop._refrescar_plan(messages, ex)
    agent_loop._refrescar_plan(messages, ex)  # dos veces: no debe acumular
    planes = [m for m in messages if m["content"].startswith(agent_loop._MARCA_PLAN)]
    assert len(planes) == 1
    assert messages[-1] is planes[0]                       # siempre al final
    cuerpo = messages[-1]["content"]
    assert "1/3 hechos" in cuerpo
    assert "[x] Investigar" in cuerpo                      # hecho
    assert "[>] Escribir el HTML" in cuerpo                # en curso
    assert "[ ] Verificar" in cuerpo                       # pendiente


def test_refrescar_plan_sin_plan_no_añade_nada(tmp_path):
    from groq_agent import agent_loop

    ex = ToolExecutor(tmp_path)
    messages = [{"role": "user", "content": "tarea"}]
    agent_loop._refrescar_plan(messages, ex)
    assert not any(m["content"].startswith(agent_loop._MARCA_PLAN) for m in messages)


def test_sembrar_manifiesto_desde_historial(tmp_path):
    from groq_agent import agent_loop

    ex = ToolExecutor(tmp_path)
    messages = [
        {"role": "tool", "content": r"Escrito C:\web\index.html (50 caracteres)."},
        {"role": "tool", "content": r"Editado C:\web\styles.css."},
        {"role": "tool", "content": r"Generado C:\docs\informe.docx. Chequeo de estructura: ok"},
    ]
    agent_loop._sembrar_manifiesto_desde_historial(messages, ex)
    assert ex.archivos_tocados[r"C:\web\index.html"] == "creado"
    assert ex.archivos_tocados[r"C:\web\styles.css"] == "editado"
    assert ex.archivos_tocados[r"C:\docs\informe.docx"] == "creado"


def test_leer_fuera_sin_aprobar_se_rechaza(tmp_path, tmp_path_factory):
    fuera = tmp_path_factory.mktemp("no_aprobada")
    (fuera / "secreto_del_usuario.txt").write_text("privado", encoding="utf-8")
    ex = ToolExecutor(tmp_path, auto_yes=False)  # sin aprobar esa carpeta
    res = ex.read_file(str(fuera / "secreto_del_usuario.txt"))
    assert res.startswith("ERROR")
    assert "privado" not in res


def test_una_web_muchos_archivos_una_sola_confirmacion(tmp_path, tmp_path_factory, monkeypatch):
    # El bug que reporto el usuario: una web son muchos ficheros en la misma
    # carpeta externa. Con 'Si, y todo en esta carpeta' se confirma UNA vez y
    # el resto se escribe sin volver a preguntar. Ademas verifica que la
    # confirmacion vive en agent_loop (hilo principal), no en la tool.
    from groq_agent import agent_loop, ui

    fuera = tmp_path_factory.mktemp("web_fuera")
    ex = ToolExecutor(tmp_path, auto_yes=False)
    llamadas = {"n": 0}

    def fake_confirm(_destino):
        llamadas["n"] += 1
        return "folder"

    monkeypatch.setattr(ui, "confirmar_ruta_externa", fake_confirm)
    monkeypatch.setattr(ui, "animate_dispatch",
                        lambda name, args, executor, root=None: executor.dispatch(name, args))
    monkeypatch.setattr(ui, "print_tool_result", lambda *a, **k: None)

    for fn, cont in [("index.html", "<h1>"), ("estilos.css", "body{}"), ("img/logo.svg", "<svg/>")]:
        destino = fuera / "casa-lucio" / fn
        r = agent_loop._confirmar_y_despachar_externo(
            "write_file", {"path": str(destino), "content": cont}, ex)
        assert r.startswith("Escrito"), r

    assert llamadas["n"] == 1  # UNA sola confirmacion para toda la carpeta
    assert (fuera / "casa-lucio" / "index.html").read_text() == "<h1>"
    assert (fuera / "casa-lucio" / "img" / "logo.svg").read_text() == "<svg/>"


def test_secreto_fuera_del_workspace_tambien_se_bloquea(tmp_path):
    from pathlib import Path

    ex = ToolExecutor(tmp_path, auto_yes=True)
    res = ex.write_file(str(Path.home() / ".env"), "SECRETO=1")
    assert res.startswith("ERROR")


def test_escribe_fuera_del_workspace_detecta(tmp_path):
    from groq_agent.tools import escribe_fuera_del_workspace

    assert escribe_fuera_del_workspace("write_file", {"path": "a.txt"}, tmp_path) is False
    assert escribe_fuera_del_workspace("write_file", {"path": "descargas/a.txt"}, tmp_path) is True
    # Una tool que no escribe nunca cuenta como "fuera".
    assert escribe_fuera_del_workspace("run_sql", {"query": "SELECT 1"}, tmp_path) is False


def test_is_system_path_bloquea_sistema_y_python(tmp_path):
    import os
    import sys
    from pathlib import Path

    from groq_agent.safety import is_system_path

    win = os.environ.get("SystemRoot") or os.environ.get("windir")
    if win:
        assert is_system_path(Path(win) / "System32" / "x.dll") is True
    # La instalacion de Python / venv activo.
    assert is_system_path(Path(sys.prefix) / "lib" / "x.py") is True


def test_is_system_path_respeta_carpetas_del_usuario(tmp_path):
    from pathlib import Path

    from groq_agent.safety import is_system_path

    home = Path.home()
    assert is_system_path(home / "Downloads" / "x.docx") is False
    assert is_system_path(home / "Desktop" / "web" / "i.html") is False
    assert is_system_path(home / "informes" / "a.txt") is False


def test_escribir_en_carpeta_de_sistema_se_bloquea_en_duro(tmp_path):
    # Ni con auto_yes: el bloqueo de sistema no lo levanta ninguna
    # confirmacion (salta ANTES de pedirla).
    import os
    from pathlib import Path

    win = os.environ.get("SystemRoot") or os.environ.get("windir") or "/etc"
    ex = ToolExecutor(tmp_path, auto_yes=True)
    res = ex.write_file(str(Path(win) / "orq_test_no_deberia.txt"), "x")
    assert res.startswith("ERROR")
    assert "SISTEMA" in res


def test_glob_and_grep(executor):
    executor.write_file("a.py", "def suma(a, b):\n    return a + b\n")
    executor.write_file("b.txt", "no es python")
    assert "a.py" in executor.glob_search("*.py")
    assert "a.py" in executor.grep_search("def suma")


def test_search_hard_cases_returns_something(executor):
    result = executor.search_hard_cases("argumento por defecto mutable", domain="python")
    assert "python-mutable-default-arg" in result


def test_generate_docx_tool(executor):
    sections = [{"heading": "Titulo", "level": 1, "text": "cuerpo"}]
    result = executor.generate_docx("out.docx", sections)
    assert "Generado" in result
    assert (executor.root / "out.docx").exists()


def test_list_available_skills_includes_python_and_office():
    skills = list_available_skills()
    assert "python-specialist" in skills
    assert "office-email-specialist" in skills
    assert "web-builder-specialist" in skills


def test_web_fetch_via_dispatch(executor):
    result = executor.dispatch("web_fetch", {"url": "https://example.com"})
    assert "Example Domain" in result


def test_lint_web_page_via_dispatch(executor):
    executor.write_file("index.html", "<html><head><title>T</title></head><body></body></html>")
    result = executor.dispatch("lint_web_page", {"path": "index.html"})
    assert "viewport" in result.lower()


def test_analyze_image_via_dispatch_on_workspace_file(executor):
    from PIL import Image

    img_path = executor.root / "bg.png"
    Image.new("RGB", (200, 100), color=(10, 10, 10)).save(img_path)
    result = executor.dispatch("analyze_image", {"url_or_path": "bg.png"})
    assert "oscura" in result


def test_search_images_via_dispatch_missing_key(executor):
    result = executor.dispatch("search_images", {"query": "pizza"})
    assert "PEXELS_API_KEY" in result


def test_render_check_via_dispatch(executor):
    executor.write_file(
        "index.html",
        "<html lang='es'><body><h1>Hola mundo</h1>"
        "<p>Una pagina de prueba con texto de sobra.</p></body></html>",
    )
    result = executor.dispatch("render_check", {"target": "index.html", "screenshot_path": "shot.png"})
    assert (executor.root / "shot_escritorio_1.png").exists()
    assert (executor.root / "shot_movil_1.png").exists()
    assert "'passed': True" in result


def test_build_system_prompt_includes_skill_and_guide():
    prompt = build_system_prompt("python-specialist")
    assert "especialista Python" in prompt
    assert "casos dificiles" in prompt.lower() or "guia de uso" in prompt.lower()


# --- formatos de numero en el .xlsx ------------------------------------
#
# Sin formato el numero es correcto pero se lee mal: un porcentaje guardado
# como fraccion (0.15, que es como Excel lo calcula) aparece "0.15" en vez
# de "15,0%". El skill de office-spreadsheet exige guardarlos asi, de modo
# que el formato no es cosmetica: es lo que hace legible el dato.


def test_el_formato_de_columna_llega_a_la_celda(tmp_path):
    from openpyxl import load_workbook

    from app.tools.xlsx_tool import create_workbook

    destino = tmp_path / "libro.xlsx"
    create_workbook(
        [{
            "name": "Datos",
            "headers": ["Concepto", "Importe", "Margen"],
            "rows": [["Ventas", 1000, 0.15]],
            "column_formats": {1: "#,##0 €", 2: "0.0%"},
        }],
        destino,
    )

    ws = load_workbook(destino)["Datos"]
    assert ws["B2"].number_format == "#,##0 €"
    assert ws["C2"].number_format == "0.0%"


def test_las_claves_de_formato_valen_como_texto(tmp_path):
    """Cuando vienen de una tool call son texto: JSON no tiene claves
    numericas. Antes se ignoraban en silencio."""
    from openpyxl import load_workbook

    from app.tools.xlsx_tool import create_workbook

    destino = tmp_path / "libro.xlsx"
    create_workbook(
        [{"name": "D", "headers": ["a", "b"], "rows": [["x", 5]],
          "column_formats": {"1": "#,##0"}}],
        destino,
    )
    assert load_workbook(destino)["D"]["B2"].number_format == "#,##0"


def test_la_celda_de_formula_hereda_el_formato_de_su_columna(tmp_path):
    """Un total de euros sin formato se lee distinto que los euros que
    suma - la tabla queda descuadrada a la vista."""
    from openpyxl import load_workbook

    from app.tools.xlsx_tool import create_workbook

    destino = tmp_path / "libro.xlsx"
    create_workbook(
        [{
            "name": "D",
            "headers": ["Concepto", "Importe"],
            "rows": [["a", 10], ["b", 20]],
            "formulas": {"B4": "=SUM(B2:B3)"},
            "column_formats": {1: "#,##0 €"},
        }],
        destino,
    )
    assert load_workbook(destino)["D"]["B4"].number_format == "#,##0 €"


def test_una_columna_de_fecha_manda_sobre_el_formato_de_columna(tmp_path):
    """Las fechas ya traen su propio formato; pisarlo con uno numerico las
    convertiria en un numero de serie a la vista."""
    from openpyxl import load_workbook

    from app.tools.xlsx_tool import create_workbook

    destino = tmp_path / "libro.xlsx"
    create_workbook(
        [{
            "name": "D",
            "headers": ["Fecha"],
            "rows": [["2026-01-15"]],
            "date_columns": [0],
            "column_formats": {0: "#,##0"},
        }],
        destino,
    )
    assert load_workbook(destino)["D"]["A2"].number_format == "DD/MM/YYYY"


def test_el_esquema_de_generate_xlsx_declara_formulas_y_formatos():
    """El fallo que esto evita: el skill exige formulas reales y el esquema
    no declaraba el campo, asi que el modelo entregaba totales calculados a
    mano sin saber que habia otra opcion."""
    from groq_agent.tools import TOOL_SCHEMAS

    esquema = next(t for t in TOOL_SCHEMAS if t["function"]["name"] == "generate_xlsx")
    hoja = esquema["function"]["parameters"]["properties"]["sheets"]["items"]["properties"]
    assert {"formulas", "date_columns", "column_formats"} <= set(hoja)


# --- tope de salida: impedir que entre, no podarlo despues -------------
#
# La poda recorta lo que YA esta en el contexto. Esto es lo otro y es mas
# barato: un `read_reference_component` de 25.000 caracteres o un
# `read_file` sin limite entraban enteros de una sentada y luego se
# reenviaban en cada vuelta. El tope vive en `dispatch`, que es por donde
# pasan todas las llamadas, para que tambien cubra las tools futuras.


def test_una_salida_larga_se_corta_en_el_dispatch(tmp_path):
    from groq_agent.tools import TOPES, ToolExecutor

    ex = ToolExecutor(tmp_path)
    ex.list_dir = lambda path=".": "z" * 50_000

    salida = ex.dispatch("list_dir", {})

    assert len(salida) < TOPES["list_dir"] + 300
    assert "cortado" in salida


def test_el_corte_dice_cuanto_falta_y_como_seguir(tmp_path):
    """Un recorte silencioso es peor que la salida larga: el modelo da por
    completo lo que no lo esta y decide sobre datos que no vio."""
    from groq_agent.tools import ToolExecutor

    ex = ToolExecutor(tmp_path)
    ex.grep_search = lambda pattern: "hallazgo\n" * 5000

    salida = ex.dispatch("grep_search", {"pattern": "x"})

    assert "de 45000 caracteres" in salida or "caracteres" in salida
    assert "afina el patron" in salida


def test_una_tool_nueva_hereda_el_tope_por_defecto(tmp_path):
    """El motivo de ponerlo en dispatch y no en cada tool: los limites
    puestos uno a uno se olvidan justo en la tool que se añade despues."""
    from groq_agent.tools import TOPE_POR_DEFECTO, ToolExecutor

    ex = ToolExecutor(tmp_path)
    ex.tool_inventada = lambda: "y" * 40_000

    salida = ex.dispatch("tool_inventada", {})

    assert len(salida) < TOPE_POR_DEFECTO + 300


def test_una_salida_corta_no_se_toca(tmp_path):
    from groq_agent.tools import ToolExecutor

    ex = ToolExecutor(tmp_path)
    ex.list_dir = lambda path=".": "tres archivos"
    assert ex.dispatch("list_dir", {}) == "tres archivos"


# --- read_file lee una ventana, no el archivo entero -------------------


def test_read_file_devuelve_una_ventana_y_dice_cuanto_queda(tmp_path):
    from groq_agent.tools import ToolExecutor

    (tmp_path / "grande.py").write_text("\n".join(f"linea {i}" for i in range(1, 1201)), encoding="utf-8")
    salida = ToolExecutor(tmp_path).read_file("grande.py")

    assert "linea 1\n" in salida
    assert "linea 401" not in salida, "no debe leer mas alla de la ventana"
    assert "quedan 800 lineas" in salida
    assert "desde=401" in salida, "tiene que decir como pedir el resto"


def test_read_file_respeta_la_ventana_pedida(tmp_path):
    from groq_agent.tools import ToolExecutor

    (tmp_path / "a.py").write_text("\n".join(f"L{i}" for i in range(1, 101)), encoding="utf-8")
    salida = ToolExecutor(tmp_path).read_file("a.py", desde=50, lineas=5)

    assert "L50" in salida and "L54" in salida
    assert "L49" not in salida and "L55" not in salida


def test_un_archivo_corto_se_devuelve_limpio(tmp_path):
    """Sin cabecera ni avisos: es el caso mas frecuente y no debe pagar
    tokens de metadatos que no aportan nada."""
    from groq_agent.tools import ToolExecutor

    (tmp_path / "corto.txt").write_text("hola\nadios", encoding="utf-8")
    assert ToolExecutor(tmp_path).read_file("corto.txt") == "hola\nadios"
