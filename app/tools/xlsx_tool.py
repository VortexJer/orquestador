"""Herramientas para el especialista office-spreadsheet: generar un
.xlsx con encabezados/formulas/fechas reales y validar la estructura -
ver skills/office-spreadsheet-specialist.md y
hard_cases/by_domain/office_spreadsheet.json.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import TypedDict

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class SheetSpec(TypedDict, total=False):
    name: str
    headers: list[str]
    rows: list[list]  # valores pueden ser str, int, float, datetime.date
    formulas: dict[str, str]  # ej. {"D2": "=B2*C2", "B10": "=SUM(B2:B9)"}
    date_columns: list[int]  # indices de columna (0-based) que son fechas
    # Formato de numero por COLUMNA (indice 0-based) - ej.
    # {1: "$#,##0", 3: "0.0%"}. Sin esto un porcentaje guardado como
    # fraccion (0.15, que es como Excel los calcula) se lee "0.15" en vez
    # de "15,0%", y una cifra de dinero sale sin separador de miles: el
    # numero es correcto pero el lector no lo puede usar.
    column_formats: dict[int, str]


def create_workbook(sheets: list[SheetSpec], output_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_spec in sheets:
        ws: Worksheet = wb.create_sheet(title=sheet_spec.get("name", "Hoja1")[:31])
        headers = sheet_spec.get("headers", [])
        rows = sheet_spec.get("rows", [])
        date_columns = {int(c) for c in sheet_spec.get("date_columns", [])}
        # Las claves llegan como texto cuando vienen de una tool call (JSON
        # no tiene claves numericas), y como int cuando se llama desde
        # Python. Sin normalizar, el formato se ignoraba en silencio: el
        # archivo salia bien salvo que los euros no eran euros.
        column_formats = {
            int(col): fmt for col, fmt in sheet_spec.get("column_formats", {}).items()
        }

        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        for row_idx, row in enumerate(rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                if (col_idx - 1) in date_columns and isinstance(value, str):
                    value = dt.datetime.strptime(value, "%Y-%m-%d").date()
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if (col_idx - 1) in date_columns:
                    cell.number_format = "DD/MM/YYYY"
                elif (col_idx - 1) in column_formats:
                    cell.number_format = column_formats[col_idx - 1]

        for cell_ref, formula in sheet_spec.get("formulas", {}).items():
            ws[cell_ref] = formula
            # Una celda de formula tambien necesita su formato: el total de
            # una columna de euros que sale sin formato se lee distinto que
            # los euros que suma.
            col_0 = ws[cell_ref].column - 1
            if col_0 in column_formats and col_0 not in date_columns:
                ws[cell_ref].number_format = column_formats[col_0]

        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def check_workbook(path: Path) -> dict:
    """Validador de estructura. Ver checklist de
    skills/office-spreadsheet-specialist.md."""
    wb = load_workbook(path)
    issues: list[str] = []

    if not wb.sheetnames:
        issues.append("El libro no tiene ninguna hoja.")

    for ws in wb.worksheets:
        if ws.max_row < 2:
            issues.append(f"Hoja '{ws.title}' no tiene filas de datos (solo encabezado o vacia).")
            continue

        header_row = [cell.value for cell in ws[1]]
        if not any(header_row):
            issues.append(f"Hoja '{ws.title}' no tiene una fila de encabezados clara en la fila 1.")

        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row <= ws.max_row and merged_range.max_row > 1:
                issues.append(
                    f"Hoja '{ws.title}' tiene celdas combinadas dentro del area de datos ({merged_range})."
                )

    return {"passed": not issues, "issues": issues}
