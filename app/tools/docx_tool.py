"""Herramientas para el especialista office-word: generar un .docx con
estructura real (estilos de Heading, listas nativas, tablas con
encabezado) y validar que esa estructura este bien formada - ver
skills/office-word-specialist.md y hard_cases/by_domain/office_word.json.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import TypedDict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.table import Table

from app.tools.image_fetch import ImagenNoValida, descargar_imagen

_PLACEHOLDER_PATTERN = re.compile(r"\[[A-ZÁÉÍÓÚÑ_ ]{2,}\]|\{\{\s*\w+\s*\}\}")
_ANCHO_IMG_DEFECTO_IN = 5.8   # cabe en el ancho util de una pagina A4/Letter


class TableSpec(TypedDict, total=False):
    headers: list[str]
    rows: list[list[str]]


class SectionSpec(TypedDict, total=False):
    heading: str
    level: int  # 1 = Heading 1, 2 = Heading 2, etc.
    text: str
    list_items: list[str]
    ordered: bool
    table: TableSpec
    image: str    # URL de una foto (de search_images) para incrustar
    caption: str  # pie de foto visible + texto alternativo (accesibilidad)


def text_to_docx(text: str, output_path: Path) -> Path:
    """Conversion simple sin estructura, para el generalist-tiny (texto
    libre de correos/resumenes). El especialista office-word usa
    create_docx en su lugar, con estructura real."""
    document = Document()
    for paragraph in text.split("\n\n"):
        document.add_paragraph(paragraph.strip())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _nivel_valido(level: object) -> int:
    """El nivel de heading puede llegar como str ('2'), float o basura: el
    modelo no siempre respeta el tipo del schema, y python-docx hace
    `0 <= level <= 9`, que revienta con 'int vs str'. Se coacciona a int en
    el rango valido (0=Titulo, 1-9=Heading), con 1 por defecto."""
    try:
        n = int(level)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return 1
    return max(0, min(9, n))


# El modelo mete markdown en el texto (**negrita**, *cursiva*) por costumbre;
# en un parrafo de Word eso sale LITERAL ("**palabra**" con los asteriscos a la
# vista). Se convierte a formato real: negrita con ** o __, cursiva con * o _.
# Los marcadores dobles van primero para que **x** no lo coma la cursiva.
_MD_INLINE = re.compile(r"\*\*(.+?)\*\*|__(.+?)__|\*(.+?)\*|_(.+?)_", re.DOTALL)


def _sin_md(texto: str) -> str:
    """Quita los marcadores markdown dejando el contenido: para sitios que ya
    tienen su propio estilo (encabezados, celdas de tabla) y no admiten runs."""
    return _MD_INLINE.sub(lambda m: m.group(1) or m.group(2) or m.group(3) or m.group(4), str(texto))


def _add_parrafo_md(document: Document, texto: str, style: str | None = None):
    """Añade un parrafo convirtiendo el markdown inline (**negrita**/*cursiva*)
    en runs con formato real, en vez de dejar los asteriscos a la vista."""
    parrafo = document.add_paragraph(style=style) if style else document.add_paragraph()
    pos = 0
    for m in _MD_INLINE.finditer(texto):
        if m.start() > pos:
            parrafo.add_run(texto[pos:m.start()])
        negrita = m.group(1) is not None or m.group(2) is not None
        contenido = m.group(1) or m.group(2) or m.group(3) or m.group(4)
        run = parrafo.add_run(contenido)
        run.bold = negrita
        run.italic = not negrita
        pos = m.end()
    if pos < len(texto):
        parrafo.add_run(texto[pos:])
    return parrafo


def _add_imagen(document: Document, url: str, caption: str | None) -> None:
    """Descarga la foto e insertarla centrada, con pie/texto-alt. Si la
    descarga falla, NO rompe el documento: deja una nota discreta y sigue -
    una foto caida no puede tumbar un informe entero."""
    try:
        data = descargar_imagen(url)
    except ImagenNoValida as exc:
        nota = document.add_paragraph(f"[imagen no disponible: {exc}]")
        for r in nota.runs:
            r.italic = True
        return
    document.add_picture(data, width=Inches(_ANCHO_IMG_DEFECTO_IN))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        # Texto alternativo para lectores de pantalla (accesibilidad).
        try:
            document.inline_shapes[-1]._inline.docPr.set("descr", str(caption))
        except Exception:  # noqa: BLE001 - el alt-text es un extra, no romper por el
            pass
        pie = document.add_paragraph()
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pie.add_run(_sin_md(str(caption)))
        run.italic = True
        run.font.size = Pt(9)


def _normalizar_seccion(section: dict) -> dict:
    """Acepta las DOS formas que puede emitir el modelo: la 'seccion' del
    schema (`heading`/`level`/`text`/`list_items`/`table`) y el 'bloque' con
    `type` que ensena el skill (`{"type":"heading"}`, `paragraph`, `list` con
    `items`, `table` con `header`). Sin esto, un bloque `{"type":"heading"}`
    entraba sin la clave `heading` -> se colaba como parrafo (documento sin
    titulos), y `list`/`table` se perdian enteros porque las claves no
    coincidian (`items` vs `list_items`, `header` vs `headers`). El resultado
    era un .docx roto aunque la tool SI se llamara."""
    tipo = section.get("type")
    if not tipo:
        return section
    if tipo == "heading":
        return {"heading": section.get("text", ""), "level": section.get("level", 1)}
    if tipo == "paragraph":
        return {"text": section.get("text", "")}
    if tipo == "list":
        return {"list_items": section.get("items", section.get("list_items", [])),
                "ordered": bool(section.get("ordered"))}
    if tipo == "table":
        headers = section.get("headers") or section.get("header") or []
        return {"table": {"headers": headers, "rows": section.get("rows", [])}}
    return section


def create_docx(sections: list[SectionSpec], output_path: Path) -> Path:
    document = Document()
    for section in (_normalizar_seccion(s) for s in sections):
        if "heading" in section:
            document.add_heading(_sin_md(section["heading"]), level=_nivel_valido(section.get("level", 1)))
        if section.get("text"):
            _add_parrafo_md(document, section["text"])
        if section.get("list_items"):
            style = "List Number" if section.get("ordered") else "List Bullet"
            for item in section["list_items"]:
                _add_parrafo_md(document, str(item), style=style)
        if section.get("table"):
            _add_table(document, section["table"])
        if section.get("image"):
            _add_imagen(document, section["image"], section.get("caption"))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _add_table(document: Document, table_spec: TableSpec) -> Table:
    headers = table_spec.get("headers", [])
    rows = table_spec.get("rows", [])
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = _sin_md(header)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = _sin_md(value)
    return table


def check_docx(path: Path) -> dict:
    """Validador de estructura (no de contenido semantico). Ver
    checklist de skills/office-word-specialist.md."""
    document = Document(path)
    issues: list[str] = []

    heading_levels = [
        int(p.style.name.replace("Heading ", ""))
        for p in document.paragraphs
        if p.style.name.startswith("Heading") and p.style.name != "Heading"
    ]

    if not heading_levels:
        issues.append("El documento no tiene ningun estilo de Heading real aplicado.")
    else:
        previous = 0
        for level in heading_levels:
            if level > previous + 1:
                issues.append(f"Salto de nivel de encabezado: de H{previous} a H{level} sin nivel intermedio.")
            previous = level

    full_text = "\n".join(p.text for p in document.paragraphs)
    placeholders = _PLACEHOLDER_PATTERN.findall(full_text)
    if placeholders:
        issues.append(f"Quedaron placeholders de plantilla sin completar: {placeholders}")

    if not full_text.strip():
        issues.append("El documento no tiene contenido de texto.")

    return {"passed": not issues, "issues": issues}
